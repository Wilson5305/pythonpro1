"""
reports.py
----------
Handles all export/reporting functionality:
  - Export student lists (View All Students) to Excel
  - Export Certificate Eligibility selections to Excel
  - Generate a printable Certificate Eligibility List (HTML -> opened in
    the default browser, where the user can use the browser's native
    Print dialog -- this is the standard, dependency-free way to offer
    "printing" from a cross-platform Tkinter desktop app).

Uses openpyxl for .xlsx generation (no external binary dependency).
"""

import os
import webbrowser
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXPORT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exports")
CERT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "certificates")

os.makedirs(EXPORT_DIR, exist_ok=True)
os.makedirs(CERT_DIR, exist_ok=True)


def _style_header(ws, num_columns):
    """Applies consistent header styling to row 1 of a worksheet."""
    header_fill = PatternFill(start_color="1F4E8C", end_color="1F4E8C", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D7DCE3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, num_columns + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border


def _autosize_columns(ws, num_columns):
    """Roughly autosizes columns based on max content length."""
    for col in range(1, num_columns + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = max(12, min(40, max_len + 4))


def export_students_to_excel(students, filename=None):
    """
    Exports a list of student rows (sqlite3.Row objects) to an Excel file.
    Returns the full path to the created file.
    """
    if not students:
        raise ValueError("No student records to export.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = [
        "Register ID", "Admission Date", "Student Name", "Qualification",
        "Phone Number", "Email Address", "Course Name", "Course Type",
        "Total Fees", "Paid Amount", "Balance Fees", "Payment Status",
        "Batch Start Date", "Faculty Name"
    ]
    ws.append(headers)
    _style_header(ws, len(headers))

    for row in students:
        ws.append([
            row["register_id"], row["admission_date"], row["student_name"],
            row["qualification"], row["phone"], row["email"], row["course_name"],
            row["course_type"], row["total_fees"], row["paid_amount"],
            row["balance_fees"], row["payment_status"],
            row["batch_start_date"] or "", row["faculty_name"] or ""
        ])

    _autosize_columns(ws, len(headers))

    if not filename:
        filename = f"Students_Export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    full_path = os.path.join(EXPORT_DIR, filename)
    wb.save(full_path)
    return full_path


def export_certificate_list_to_excel(students, filename=None):
    """
    Exports the selected Certificate Eligibility list to Excel.
    `students` is a list of sqlite3.Row (or dict-like) objects.
    """
    if not students:
        raise ValueError("No students selected for export.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Certificate Eligibility"

    headers = ["Register ID", "Student Name", "Course Name",
               "Admission Date", "Completion Status"]
    ws.append(headers)
    _style_header(ws, len(headers))

    for row in students:
        ws.append([
            row["register_id"], row["student_name"], row["course_name"],
            row["admission_date"], "Eligible - Fully Paid"
        ])

    _autosize_columns(ws, len(headers))

    if not filename:
        filename = f"Certificate_Eligibility_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    full_path = os.path.join(EXPORT_DIR, filename)
    wb.save(full_path)
    return full_path


def generate_printable_certificate_list(students):
    """
    Generates a clean, print-friendly HTML page listing the selected
    certificate-eligible students, then opens it in the system's default
    web browser. The user can then use Ctrl+P / the browser Print dialog,
    which is the standard cross-platform way to print from a Tkinter app.

    Returns the path to the generated HTML file.
    """
    if not students:
        raise ValueError("No students selected to print.")

    rows_html = ""
    for idx, row in enumerate(students, start=1):
        rows_html += f"""
            <tr>
                <td>{idx}</td>
                <td>{row['register_id']}</td>
                <td>{row['student_name']}</td>
                <td>{row['course_name']}</td>
                <td>{row['admission_date']}</td>
                <td class="status">Eligible</td>
            </tr>
        """

    generated_on = datetime.now().strftime("%d-%b-%Y %I:%M %p")

    html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <title>Certificate Eligibility List</title>
        <style>
            body {{ font-family: 'Segoe UI', Arial, sans-serif; margin: 40px; color: #1E2A38; }}
            h1 {{ color: #1F4E8C; border-bottom: 3px solid #1F4E8C; padding-bottom: 10px; }}
            .meta {{ color: #6B7785; margin-bottom: 20px; font-size: 13px; }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 10px; }}
            th, td {{ border: 1px solid #D7DCE3; padding: 8px 12px; text-align: left; }}
            th {{ background-color: #1F4E8C; color: white; }}
            tr:nth-child(even) {{ background-color: #F0F3F7; }}
            .status {{ color: #2E9E6B; font-weight: bold; }}
            @media print {{
                button {{ display: none; }}
            }}
        </style>
    </head>
    <body>
        <h1>Course Completion Certificate Eligibility List</h1>
        <div class="meta">Generated on: {generated_on} &nbsp;|&nbsp; Total Eligible Students: {len(students)}</div>
        <table>
            <thead>
                <tr>
                    <th>#</th>
                    <th>Register ID</th>
                    <th>Student Name</th>
                    <th>Course Name</th>
                    <th>Admission Date</th>
                    <th>Status</th>
                </tr>
            </thead>
            <tbody>
                {rows_html}
            </tbody>
        </table>
        <br>
        <button onclick="window.print()" style="padding:10px 20px; background:#1F4E8C; color:white; border:none; border-radius:4px; cursor:pointer; font-size:14px;">
            Print this List
        </button>
    </body>
    </html>
    """

    filename = f"Certificate_List_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
    full_path = os.path.join(CERT_DIR, filename)
    with open(full_path, "w", encoding="utf-8") as f:
        f.write(html)

    webbrowser.open(f"file://{full_path}")
    return full_path
